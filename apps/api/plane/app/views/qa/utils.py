import os
import json
import random
import re
import string
import uuid
from io import BytesIO

from django.core.files.uploadedfile import InMemoryUploadedFile
from django.db.models import Q
from openpyxl import load_workbook

from plane.db.models import CaseModule


def expand_module_subtree_ids(module_id):
    """把模块 id 展开为含全部后代模块的 id 列表（BFS），与模块树计数口径一致。

    供用例列表按模块过滤、导出、模板用例列表共用。
    """
    expanded = {str(module_id)}
    frontier = [str(module_id)]
    while frontier:
        children = list(
            CaseModule.objects.filter(
                parent_id__in=frontier, deleted_at__isnull=True
            ).values_list("id", flat=True)
        )
        new_children = [str(c) for c in children if str(c) not in expanded]
        if not new_children:
            break
        expanded.update(new_children)
        frontier = new_children
    return list(expanded)


# 模块树「未归类」节点 / 模块筛选「未归类」选项在 module_id 参数里的取值：module 为空的条目
UNASSIGNED_MODULE_ID = "none"


def expand_subtree_ids(model, module_ids):
    """任意带 parent 自关联的模块模型：把一组模块 id 展开为含全部后代的 id 列表（BFS）。"""
    expanded = {str(mid) for mid in module_ids}
    frontier = list(expanded)
    while frontier:
        children = model.objects.filter(
            parent_id__in=frontier, deleted_at__isnull=True
        ).values_list("id", flat=True)
        new_children = [str(c) for c in children if str(c) not in expanded]
        if not new_children:
            break
        expanded.update(new_children)
        frontier = new_children
    return list(expanded)


def build_module_filter_q(raw, model=CaseModule, field="module_id"):
    """按模块过滤的 Q：真实模块递归包含子模块，"none" 表示 field 为空，两者取或。

    raw 接受单值、逗号分隔串或列表；非法 id 忽略。没有任何有效值时返回 None（调用方不过滤）。
    """
    values = raw if isinstance(raw, (list, tuple)) else [raw]
    module_ids = []
    include_unassigned = False
    for value in values:
        for part in str(value or "").split(","):
            part = part.strip()
            if not part:
                continue
            if part == UNASSIGNED_MODULE_ID:
                include_unassigned = True
                continue
            try:
                module_ids.append(str(uuid.UUID(part)))
            except ValueError:
                continue
    if not module_ids and not include_unassigned:
        return None
    q = Q()
    if module_ids:
        q |= Q(**{f"{field}__in": expand_subtree_ids(model, module_ids)})
    if include_unassigned:
        q |= Q(**{f"{field}__isnull": True})
    return q


def build_case_activity_snapshot(case):
    """更新前的用例快照（JSON 字符串），供活动流 diff 使用。

    与 CaseAPIView.put / TemplateCaseAPIView.put 共用。
    """
    return json.dumps({
        "name": case.name,
        "type": case.type,
        "test_type": case.test_type,
        "priority": case.priority,
        "assignee_id": str(case.assignee_id) if case.assignee_id else None,
        "module_id": str(case.module_id) if case.module_id else None,
        "labels": [str(l) for l in case.labels.values_list("id", flat=True)],
        "precondition": case.precondition,
        "steps": case.steps,
        "text_description": case.text_description,
        "text_result": case.text_result,
        "remark": case.remark,
    })


def split_by_numbering(text):
    # 按照 "1. xxx", "2. xxx" 分割文本，但只匹配纯数字+点开头的行首
    pattern = r'(\d+\.)\s*(.*?)(?=\n\d+\.|\Z)'
    matches = re.findall(pattern, text.strip(), re.DOTALL)
    return {int(num[:-1]): content.strip() for num, content in matches}


def build_description_result_list(description: str, result: str):
    desc_dict = split_by_numbering(description)
    res_dict = split_by_numbering(result)

    result_list = []
    # 遍历所有出现在 description 中的编号
    for key in sorted(desc_dict.keys()):
        description_text = desc_dict[key]
        result_text = res_dict.get(key, "")  # 若无对应结果则返回空字符串
        result_list.append({
            "description": description_text,
            "result": result_text
        })

    return result_list


def parser_excel_case(file_path, sheet_name='case'):
    """使用openpyxl将Excel转换为字典"""

    mapping = {"功能": 'label', '测试内容': 'name', '用例等级': 'priority', '测试目的': 'remark',
               '预置条件': 'precondition', '测试步骤': 'description', '预期结果': 'result','模块':'module'}

    # 加载工作簿
    workbook = load_workbook(file_path)

    # 选择工作表
    if sheet_name:
        worksheet = workbook[sheet_name]
    else:
        worksheet = workbook.active  # 默认活动工作表

    # 获取第一行作为列标题
    headers = [(mapping.get(cell.value) or cell.value) for cell in worksheet[1]]

    # 读取数据行
    data = []
    for row in worksheet.iter_rows(min_row=2, values_only=True):
        # 创建字典，跳过空行
        if any(cell is not None for cell in row):
            row_dict = dict(zip(headers, row))
            data.append({key: value for key, value in row_dict.items() if key in mapping.values()})

    return data


def get_extension_without_dot(filename):
    """获取不带点号的扩展名"""
    _, extension = os.path.splitext(filename)
    return extension[1:].lower() if extension else ''


def clear_excel_data(excel_data: list[dict]):
    excel = []
    for data in excel_data:
        description = data.pop('description')
        result = data.pop('result')
        steps = build_description_result_list(description, result)
        data['steps'] = steps
        # 标签
        data['label'] = [label.strip() for label in data['label'].split('\n') if label]

        excel.append(data)
    return excel


def parser_case_file(files: list[InMemoryUploadedFile]) -> list:
    result = list()
    for file in files:
        if (suffix := get_extension_without_dot(file.name)) in ['json']:
            result.extend(json.load(file))
        elif suffix in ['xlsx']:
            data = parser_excel_case(BytesIO(file.read()))
            result.extend(clear_excel_data(data))
        else:
            raise Exception('不是支持的文件类型')
    return result
